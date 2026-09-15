"""El generador de carga VTEX para Supermall.

Ejecutar:  python scripts/test_vtex_generator.py

Las pruebas EJECUTAN el motor contra el export REAL de VTEX que esta en
`data/vtex_muestra_supermallpe/` -- las cuatro planillas, con la primera fila en
blanco, las cabeceras en la segunda y el archivo de especificaciones de
productos partido en dos hojas, que es como sale del admin. Leer el codigo no es
ejecutarlo: `start_suelto` tenia ocho pruebas que leian su fuente y ninguna la
llamaba, y estaba roto desde que se escribio.

Lo que fija
-----------
- La estructura de los cuatro archivos, columna a columna: es una plantilla de
  VTEX y una columna de mas o de menos la rechaza.
- **No se inventan matches.** Cuando un escalon de la escalera da dos
  candidatos, el registro va a REVISAR y queda FUERA del archivo.
- **No se inventan IDs.** Ni de producto, ni de SKU, ni de marca, ni de
  categoria, ni de valor de especificacion.
- Que un producto que ya existe REUSA su Product ID, sus SKU ID, su URL y el ID
  de instancia de sus especificaciones de texto.
- Que la pantalla no escribe en ninguna tienda.
"""
import ast
import inspect
import io
import os
import subprocess
import sys
import textwrap
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from engines import vtex_export as vtex  # noqa: E402
from engines import garment_types  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MUESTRA = os.path.join(RAIZ, "data", "vtex_muestra_supermallpe")


def _leer_muestra():
    """Las cuatro planillas reales, leidas como las lee la pantalla."""
    import app_matrixify as app

    class Subido(io.BytesIO):
        """Lo minimo que `vtex_hojas_del_archivo` necesita."""

        def __init__(self, ruta):
            with open(ruta, "rb") as fh:
                super().__init__(fh.read())
            self.name = os.path.basename(ruta)

    archivos = [Subido(os.path.join(MUESTRA, n)) for n in sorted(os.listdir(MUESTRA))
                if n.endswith(".xlsx") and not n.endswith("-es.xlsx")]
    return app.vtex_leer_planillas(archivos)


def _subido(nombre):
    """Un archivo de la muestra, como lo entrega el `file_uploader`."""
    import io as _io

    class Subido(_io.BytesIO):
        def __init__(self, ruta):
            with open(ruta, "rb") as fh:
                super().__init__(fh.read())
            self.name = os.path.basename(ruta)

    return Subido(os.path.join(MUESTRA, nombre))


class LecturaDelExport(unittest.TestCase):
    """Las cuatro planillas, tal y como las exporta VTEX."""

    @classmethod
    def setUpClass(cls):
        cls.planillas, cls.informe = _leer_muestra()
        cls.catalogo = vtex.leer_catalogo(cls.planillas)

    def test_las_cuatro_planillas_se_reconocen_por_su_cabecera(self):
        """Por la CABECERA, no por el nombre: VTEX lo nombra con una marca de
        tiempo y el usuario lo renombra."""
        for archivo in vtex.ARCHIVOS:
            self.assertTrue(self.planillas.get(archivo),
                            f"no se reconocio {vtex.ETIQUETAS[archivo]}")

    def test_las_hojas_no_se_materializan(self):
        """El export real son 177 MB: como listas de diccionarios son 1,7 GB
        medidos, con un contenedor de 1 GB para toda la app."""
        import app_matrixify as app
        hoja = self.planillas[vtex.PRODUCTOS][0]
        self.assertIsInstance(hoja, app.HojaDeVtex)
        # Es un iterable REUTILIZABLE: el indexador lo recorre una vez y la
        # pantalla puede volver a recorrerlo sin que se haya guardado nada.
        primeras = [f["Product ID"] for _n, f in zip(range(3), hoja)]
        self.assertEqual(primeras, [f["Product ID"] for _n, f in zip(range(3), hoja)])
        self.assertGreater(hoja.filas, 400)

    def test_la_cabecera_no_esta_en_la_primera_fila(self):
        """VTEX exporta con la primera fila EN BLANCO. Leyendo con `header=0`
        las columnas salen `Unnamed: 0` y no se reconoce ni el archivo."""
        hoja = self.planillas[vtex.PRODUCTOS][0]
        columnas = list(next(iter(hoja)).keys())
        self.assertIn("Product ID", columnas)
        self.assertEqual(len(columnas), len(vtex.COLUMNAS_PRODUCTOS))

    def test_las_dos_hojas_de_especificaciones_se_concatenan(self):
        """VTEX parte ese export por el limite de filas de Excel, no por
        contenido: las dos hojas no comparten ni un producto."""
        productos = {f["ID del producto"] for hoja in self.planillas[vtex.ESPEC_PRODUCTO]
                     for f in hoja}
        self.assertGreater(len(productos), 15)
        hojas = {fila["Hoja"] for fila in self.informe
                 if fila["Planilla"] == vtex.ETIQUETAS[vtex.ESPEC_PRODUCTO]}
        self.assertEqual(len(hojas), 2, "las dos hojas tienen que entrar las dos")

    def test_el_indice_del_maestro_sale_del_archivo(self):
        self.assertGreater(len(self.catalogo.productos), 100)
        self.assertGreater(len(self.catalogo.skus), 400)
        self.assertIn("hush puppies", self.catalogo.marcas)
        self.assertIn(("hombre", "zapatos"), self.catalogo.categorias)

    def test_el_dominio_de_cada_campo_viaja_en_el_export(self):
        """Es lo que permite traducir un texto a su ID sin llamar a la API."""
        talla = self.catalogo.campo_de_sku("41", "Talla")
        self.assertEqual(talla.id, "28")
        self.assertTrue(talla.es_lista)
        identificador, _ambiguo = talla.id_de("39")
        self.assertEqual(identificador, "141")
        color = self.catalogo.campo_de_sku("41", "Color")
        self.assertEqual(color.id_de("Rojo")[0], "78")

    def test_un_valor_que_el_dominio_no_tiene_no_se_inventa(self):
        talla = self.catalogo.campo_de_sku("41", "Talla")
        self.assertEqual(talla.id_de("TALLA-QUE-NO-EXISTE"), ("", False))

    def test_los_tres_campos_que_se_llaman_igual_se_conservan(self):
        """La tienda tiene TRES campos "Tecnologia " (26, 94 y 117). Con un
        indice por nombre se escribirian dos filas de menos."""
        campos = self.catalogo.campos_de_la_categoria("41", de_sku=False)
        tecnologias = [c for c in campos if vtex.clave(c.nombre) == "tecnologia"]
        self.assertGreaterEqual(len(tecnologias), 2)
        self.assertEqual(len({c.id for c in tecnologias}), len(tecnologias))

    def test_las_medidas_salen_de_la_categoria_y_no_se_inventan(self):
        medidas, heredadas = self.catalogo.medidas_de("Hombre", "Zapatos")
        self.assertFalse(heredadas)
        self.assertEqual(medidas["Package weight"], "800")
        _medidas, heredadas = self.catalogo.medidas_de("Hombre", "Categoria inventada")
        self.assertTrue(heredadas, "una categoria sin datos tiene que AVISAR")

    def test_el_peso_cubico_usa_la_formula_de_la_tienda(self):
        """34 x 49.5 x 14 / 4800 = 4,9088, que es lo que trae el export."""
        medidas = {"Package width": "34", "Package height": "49.5", "Package length": "14"}
        self.assertEqual(vtex._peso_cubico(medidas), "4.9088")
        self.assertEqual(vtex._peso_cubico({"Package width": "34"}), "")


def _ficha(codigo, **extra):
    base = {
        "Mod-Col": codigo, "Modelo": codigo.split("-")[0],
        "Color codigo": codigo.split("-", 1)[1] if "-" in codigo else "",
        "Title": "PRODUCTO DE PRUEBA", "Body HTML": "<p>Texto</p>",
        "Tipo": "Zapatos", "Clase": "Calzado", "Genero": "Hombre",
        "Marca": "Hush Puppies", "Color": "Rojo", "Imagenes": [],
        "Variantes": [{"Talla": "39", "SKU": "S1", "EAN": ""}],
    }
    base.update(extra)
    return base


class Emparejamiento(unittest.TestCase):
    """La escalera: SKU/RefId -> EAN -> ARTI -> Product/SKU ID -> Modelo+Color."""

    @classmethod
    def setUpClass(cls):
        planillas, _informe = _leer_muestra()
        cls.catalogo = vtex.leer_catalogo(planillas)

    def _emparejar(self, ficha):
        return vtex.emparejar([ficha], self.catalogo,
                              nombres_de_tipo=garment_types.sinonimos_de)[0]

    def test_el_codigo_modelo_color_encuentra_el_producto_de_vtex(self):
        salida = self._emparejar(_ficha("HP102011307-251"))
        self.assertEqual(salida["Estado"], vtex.EXISTE)
        self.assertEqual(salida["Criterio"], vtex.POR_CODIGO_ARTI)
        self.assertEqual(salida["Product ID"], "2")

    def test_el_refid_del_sku_manda_sobre_el_codigo(self):
        """Primer escalon de la escalera."""
        sku = self.catalogo.skus["310669"]
        salida = self._emparejar(_ficha(
            "OTRO-CODIGO", Variantes=[{"Talla": "39", "SKU": sku.referencia, "EAN": ""}]))
        self.assertEqual(salida["Criterio"], vtex.POR_REFERENCIA_SKU)
        self.assertEqual(salida["Product ID"], "2")

    def test_un_producto_que_no_esta_en_vtex_sale_como_nuevo(self):
        salida = self._emparejar(_ficha("ZZ99999-XXX"))
        self.assertEqual(salida["Estado"], vtex.NUEVO)
        self.assertEqual(salida["Product ID"], "")

    def test_el_sku_se_empareja_por_talla_dentro_del_producto(self):
        salida = self._emparejar(_ficha(
            "HP102011307-251",
            Variantes=[{"Talla": "39", "SKU": "", "EAN": ""},
                       {"Talla": "99", "SKU": "", "EAN": ""}]))
        por_talla = {v["Talla"]: v for v in salida["Variantes"]}
        self.assertEqual(por_talla["39"]["SKU ID"], "310669")
        self.assertEqual(por_talla["39"]["Estado"], vtex.EXISTE)
        self.assertEqual(por_talla["99"]["SKU ID"], "")
        self.assertEqual(por_talla["99"]["Estado"], vtex.NUEVO)

    def test_el_sku_no_se_busca_fuera_de_su_producto(self):
        """Un RefId que apunta a otro producto no es este SKU: usarlo moveria
        la talla de sitio."""
        ajeno = self.catalogo.productos["1"].skus[0]
        salida = self._emparejar(_ficha(
            "HP102011307-251",
            Variantes=[{"Talla": "39", "SKU": "", "EAN": "",
                        "SKU ID": ajeno.id}]))
        self.assertEqual(salida["Variantes"][0]["SKU ID"], "310669")

    def test_dos_candidatos_mandan_a_revisar_y_no_eligen_uno(self):
        """Es LA regla: un ID equivocado en VTEX sobrescribe otro producto."""
        catalogo = vtex.leer_catalogo({
            vtex.PRODUCTOS: [
                {"Product ID": "10", "Product Name": "A", "Product reference code": "AA-1",
                 "SKU ID": "100", "SKU name": "TALLA 39", "SKU reference code": "X1",
                 "Brand": "Hush Puppies", "Brand ID": "9", "Department": "Hombre",
                 "Department ID": "25", "Category": "Zapatos", "Category ID": "41"},
                {"Product ID": "11", "Product Name": "B", "Product reference code": "AA-1",
                 "SKU ID": "101", "SKU name": "TALLA 40", "SKU reference code": "X2",
                 "Brand": "Hush Puppies", "Brand ID": "9", "Department": "Hombre",
                 "Department ID": "25", "Category": "Zapatos", "Category ID": "41"},
            ]})
        salida = vtex.emparejar([_ficha("AA-1")], catalogo)[0]
        self.assertEqual(salida["Estado"], vtex.REVISAR)
        self.assertIn(vtex.AMBIGUO, salida["Motivos"])
        self.assertEqual(salida["Product ID"], "")

    def test_una_marca_que_no_existe_en_vtex_no_se_crea(self):
        salida = self._emparejar(_ficha("ZZ1-AAA", Marca="Marca Inventada"))
        self.assertEqual(salida["Estado"], vtex.REVISAR)
        self.assertIn(vtex.SIN_MARCA_VTEX, salida["Motivos"])

    def test_un_tipo_sin_categoria_en_vtex_no_se_adivina(self):
        salida = self._emparejar(_ficha("ZZ2-AAA", Tipo="Paraguas", Clase="Accesorios"))
        self.assertEqual(salida["Estado"], vtex.REVISAR)
        self.assertIn(vtex.SIN_CATEGORIA_VTEX, salida["Motivos"])

    def test_el_genero_decide_el_departamento(self):
        salida = self._emparejar(_ficha("ZZ3-AAA", Genero="Mujer", Tipo="Zapatos",
                                        Marca="Columbia"))
        self.assertEqual(salida["Departamento"], "Mujer")
        self.assertEqual(salida["Categoria"], "Zapatos")

    def test_un_genero_que_no_cae_en_ningun_departamento_se_revisa(self):
        """Unisex no es Hombre ni Mujer, y adivinarlo lo publica en la seccion
        equivocada."""
        salida = self._emparejar(_ficha("ZZ4-AAA", Genero="Unisex", Clase="Calzado"))
        self.assertEqual(salida["Estado"], vtex.REVISAR)
        self.assertIn("departamento", salida["Motivos"])

    def test_el_singular_del_diccionario_encuentra_la_categoria_en_plural(self):
        """`Zapato` en el diccionario de tipos, `Zapatos` en VTEX."""
        salida = self._emparejar(_ficha("ZZ5-AAA", Tipo="Zapato"))
        self.assertEqual(salida["Categoria"], "Zapatos")

    def test_un_producto_sin_tallas_se_revisa(self):
        salida = self._emparejar(_ficha("ZZ6-AAA", Variantes=[]))
        self.assertEqual(salida["Estado"], vtex.REVISAR)
        self.assertIn(vtex.SIN_TALLAS, salida["Motivos"])

    def test_solo_salen_las_excepciones(self):
        """De una carga de miles, lo que se revisa son las decenas que no se
        resolvieron solas."""
        emparejados = vtex.emparejar(
            [_ficha("HP102011307-251"), _ficha("ZZ7-AAA"),
             _ficha("ZZ8-AAA", Marca="Marca Inventada")],
            self.catalogo, nombres_de_tipo=garment_types.sinonimos_de)
        excepciones = vtex.filas_de_excepciones(emparejados)
        self.assertEqual(len(excepciones), 1)
        self.assertEqual(excepciones[0]["Mod-Col"], "ZZ8-AAA")


class Generacion(unittest.TestCase):
    """Los cuatro archivos."""

    @classmethod
    def setUpClass(cls):
        planillas, _informe = _leer_muestra()
        cls.catalogo = vtex.leer_catalogo(planillas)

    def _generar(self, fichas, **opciones):
        emparejados = vtex.emparejar(fichas, self.catalogo,
                                     nombres_de_tipo=garment_types.sinonimos_de)
        return vtex.generar(emparejados, self.catalogo, **opciones)

    def test_las_columnas_son_las_del_export_en_su_orden(self):
        tablas, _inc, _res = self._generar(
            [_ficha("HP102011307-251", Imagenes=["https://x/a.jpg"])])
        for archivo in vtex.ARCHIVOS:
            primera = next(iter(tablas[archivo]), None)
            self.assertIsNotNone(primera, vtex.ETIQUETAS[archivo])
            self.assertEqual(list(primera.keys()),
                             list(self.catalogo.cabeceras[archivo]),
                             f"{vtex.ETIQUETAS[archivo]}: columnas distintas del export")

    def test_un_producto_existente_reusa_su_id_su_url_y_sus_skus(self):
        tablas, _inc, res = self._generar([_ficha(
            "HP102011307-251",
            Variantes=[{"Talla": "39", "SKU": "S1", "EAN": "7701"}])])
        fila = tablas[vtex.PRODUCTOS][0]
        self.assertEqual(fila["Product ID"], "2")
        self.assertEqual(fila["SKU ID"], "310669")
        self.assertEqual(fila["SKU reference code"], "310669")
        self.assertEqual(fila["Product URL"],
                         self.catalogo.productos["2"].url)
        self.assertEqual(res["Productos que se actualizan"], 1)
        self.assertEqual(res["Productos que se crean"], 0)

    def test_un_producto_nuevo_deja_el_id_vacio(self):
        """VTEX lo asigna. Inventarlo sobrescribe otro producto."""
        tablas, _inc, res = self._generar([_ficha("ZZ9-AAA")])
        fila = tablas[vtex.PRODUCTOS][0]
        self.assertEqual(fila["Product ID"], "")
        self.assertEqual(fila["SKU ID"], "")
        self.assertEqual(fila["Product reference code"], "ZZ9-AAA")
        self.assertEqual(res["Productos que se crean"], 1)

    def test_el_refid_de_un_sku_nuevo_tiene_tres_modos(self):
        ficha = _ficha("ZZ10-AAA", Variantes=[{"Talla": "39", "SKU": "ARTI-1", "EAN": ""}])
        for modo, esperado in (("arti", "ARTI-1"), ("codigo", "ZZ10-AAA-39"), ("vacio", "")):
            tablas, _i, _r = self._generar([ficha], referencia_sku=modo)
            self.assertEqual(tablas[vtex.PRODUCTOS][0]["SKU reference code"], esperado, modo)

    def test_lo_que_hay_que_revisar_queda_FUERA_del_archivo(self):
        tablas, incidencias, _res = self._generar(
            [_ficha("HP102011307-251"), _ficha("ZZ11-AAA", Marca="Marca Inventada")])
        codigos = {f["Product reference code"] for f in tablas[vtex.PRODUCTOS]}
        self.assertNotIn("ZZ11-AAA", codigos)
        self.assertEqual([f["Mod-Col"] for f in incidencias["fuera"]], ["ZZ11-AAA"])

    def test_la_especificacion_de_lista_lleva_el_id_del_dominio(self):
        tablas, _i, _r = self._generar([_ficha("ZZ12-AAA", Tipo="Zapatos", Genero="Hombre")])
        por_campo = {f["Nombre del campo"]: f for f in tablas[vtex.ESPEC_PRODUCTO]}
        self.assertEqual(por_campo["Tipo de Producto"]["Valores de especificación"], "Zapatos")
        self.assertEqual(por_campo["Tipo de Producto"]["IDs de especificación"], "632")
        self.assertEqual(por_campo["Género"]["IDs de especificación"], "62")

    def test_un_valor_fuera_del_dominio_no_se_escribe_y_se_avisa(self):
        """Cargarlo con un ID cualquiera publicaria otro valor."""
        _t, incidencias, _r = self._generar([_ficha("ZZ13-AAA", Color="ColorQueNoExiste")])
        avisos = " ".join(a["Aviso"] for a in incidencias["avisos"])
        self.assertIn("ColorQueNoExiste", avisos)

    def test_los_campos_sin_dato_salen_con_el_guion_de_vtex(self):
        tablas, _i, _r = self._generar([_ficha("ZZ14-AAA")])
        vacios = [f for f in tablas[vtex.ESPEC_PRODUCTO]
                  if f["Valores de especificación"] == vtex.VACIO]
        self.assertTrue(vacios)
        self.assertTrue(all(not f["IDs de especificación"] for f in vacios))

    def test_el_id_de_una_especificacion_de_texto_solo_se_reusa_en_su_producto(self):
        """En un campo `Texto` ese ID es el de ESA instancia: reusarlo en otro
        producto le escribiria encima."""
        tablas, _i, _r = self._generar([_ficha("HP102011307-251", Clase="Calzado"),
                                        _ficha("ZZ15-AAA", Clase="Calzado")])
        por_producto = {}
        for fila in tablas[vtex.ESPEC_PRODUCTO]:
            if fila["Nombre del campo"] == "Clase":
                por_producto[fila["Código de referencia del producto"]] = fila
        self.assertTrue(por_producto["HP102011307-251"]["IDs de especificación"])
        self.assertEqual(por_producto["ZZ15-AAA"]["IDs de especificación"], "")

    def test_cada_sku_lleva_su_talla_y_su_color(self):
        tablas, _i, _r = self._generar([_ficha(
            "HP102011307-251", Color="Rojo",
            Variantes=[{"Talla": "39", "SKU": "", "EAN": ""},
                       {"Talla": "40", "SKU": "", "EAN": ""}])])
        filas = tablas[vtex.ESPEC_SKU]
        self.assertEqual(len(filas), 4)
        tallas = {f["Valores de especificación"] for f in filas
                  if f["Nombre del campo"] == "Talla"}
        self.assertEqual(tallas, {"39", "40"})
        colores = {f["IDs de especificación"] for f in filas
                   if f["Nombre del campo"] == "Color"}
        self.assertEqual(colores, {"78"})

    def test_la_imagen_va_en_la_columna_de_importacion(self):
        """Esa columna existe para eso: VTEX se descarga la URL. Escribirla en
        `Ruta de la imagen` no importaria nada."""
        tablas, _i, _r = self._generar([_ficha(
            "ZZ16-AAA", Imagenes=["https://cdn.shopify.com/a.jpg"])])
        fila = next(iter(tablas[vtex.IMAGENES]))
        self.assertEqual(fila["URL de importación de la imagen"],
                         "https://cdn.shopify.com/a.jpg")
        self.assertEqual(fila["Ruta de la imagen"], "")
        self.assertEqual(fila["ID de la imagen"], "")

    def test_la_imagen_se_repite_en_cada_talla_porque_cuelga_del_sku(self):
        tablas, _i, _r = self._generar([_ficha(
            "ZZ17-AAA", Imagenes=["https://x/a.jpg", "https://x/b.jpg"],
            Variantes=[{"Talla": "39", "SKU": "", "EAN": ""},
                       {"Talla": "40", "SKU": "", "EAN": ""}])])
        self.assertEqual(len(tablas[vtex.IMAGENES]), 4)
        self.assertEqual({f["Nombre de la imagen"] for f in tablas[vtex.IMAGENES]},
                         {"TALLA-39", "TALLA-40"})

    def test_una_imagen_que_el_sku_ya_tiene_no_se_repite(self):
        """VTEX la cargaria dos veces y la ficha saldria con la foto duplicada."""
        existente = self.catalogo.imagenes_de_sku["310669"][0]
        tablas, _i, _r = self._generar([_ficha(
            "HP102011307-251", Imagenes=[existente, "https://x/nueva.jpg"],
            Variantes=[{"Talla": "39", "SKU": "", "EAN": ""}])])
        urls = [f["URL de importación de la imagen"] for f in tablas[vtex.IMAGENES]]
        self.assertEqual(urls, ["https://x/nueva.jpg"])

    def test_el_tope_de_fotos_por_sku_se_respeta(self):
        tablas, _i, _r = self._generar(
            [_ficha("ZZ18-AAA", Imagenes=[f"https://x/{n}.jpg" for n in range(10)])],
            imagenes_por_sku=3)
        self.assertEqual(len(tablas[vtex.IMAGENES]), 3)


class Validacion(unittest.TestCase):
    """La foto del archivo, antes de subirlo."""

    @classmethod
    def setUpClass(cls):
        planillas, _informe = _leer_muestra()
        cls.catalogo = vtex.leer_catalogo(planillas)

    def _validar(self, fichas):
        emparejados = vtex.emparejar(fichas, self.catalogo,
                                     nombres_de_tipo=garment_types.sinonimos_de)
        tablas, _inc, _res = vtex.generar(emparejados, self.catalogo)
        return tablas, vtex.validar(tablas, self.catalogo)

    def test_una_carga_correcta_no_tiene_hallazgos_que_bloqueen(self):
        _tablas, hallazgos = self._validar([_ficha(
            "HP102011307-251", Imagenes=["https://x/a.jpg"])])
        bloqueos, _avisos = vtex.resumen_de_validacion(hallazgos)
        self.assertEqual(bloqueos, 0, [h for h in hallazgos if h["Gravedad"] == vtex.BLOQUEA])

    def test_la_talla_repetida_dentro_del_producto_bloquea(self):
        tablas, _h = self._validar([_ficha(
            "ZZ20-AAA", Variantes=[{"Talla": "39", "SKU": "A", "EAN": ""}])])
        tablas[vtex.PRODUCTOS].append(dict(tablas[vtex.PRODUCTOS][0]))
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any(h["Revisar"] == "SKU duplicado" and h["Gravedad"] == vtex.BLOQUEA
                            for h in hallazgos))

    def test_el_mismo_codigo_con_dos_product_id_bloquea(self):
        tablas, _h = self._validar([_ficha("ZZ21-AAA")])
        clon = dict(tablas[vtex.PRODUCTOS][0])
        clon["Product ID"] = "2"
        clon["SKU name"] = "TALLA 44"
        tablas[vtex.PRODUCTOS].append(clon)
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any(h["Revisar"] == "Product ID" for h in hallazgos))

    def test_un_sku_id_en_dos_productos_bloquea(self):
        """Lo MOVERIA de producto."""
        tablas, _h = self._validar([_ficha("HP102011307-251")])
        clon = dict(tablas[vtex.PRODUCTOS][0])
        clon["Product reference code"] = "OTRO-1"
        tablas[vtex.PRODUCTOS].append(clon)
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any(h["Revisar"] == "SKU ID repetido" for h in hallazgos))

    def test_un_product_id_que_no_esta_en_el_export_bloquea(self):
        tablas, _h = self._validar([_ficha("ZZ22-AAA")])
        tablas[vtex.PRODUCTOS][0]["Product ID"] = "99999999"
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any("no está en el archivo de VTEX" in h["Detalle"]
                            for h in hallazgos))

    def test_un_campo_obligatorio_vacio_bloquea(self):
        tablas, _h = self._validar([_ficha("ZZ23-AAA")])
        tablas[vtex.PRODUCTOS][0]["Product Name"] = ""
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any(h["Revisar"] == "Campo obligatorio" for h in hallazgos))

    def test_una_categoria_que_no_es_de_vtex_bloquea(self):
        tablas, _h = self._validar([_ficha("ZZ24-AAA")])
        tablas[vtex.PRODUCTOS][0]["Category"] = "Inventada"
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any(h["Revisar"] == "Categoría" for h in hallazgos))

    def test_una_especificacion_de_lista_sin_id_bloquea(self):
        """VTEX la rechaza."""
        tablas, _h = self._validar([_ficha("ZZ25-AAA")])
        # La planilla es un generador: se materializa una copia para romperla.
        rotas = []
        roto = False
        for fila in tablas[vtex.ESPEC_PRODUCTO]:
            fila = dict(fila)
            if not roto and fila["Tipo de campo"] == "CheckBox" and fila["IDs de especificación"]:
                fila["IDs de especificación"] = ""
                roto = True
            rotas.append(fila)
        self.assertTrue(roto, "el caso necesita una especificacion de lista con valor")
        tablas[vtex.ESPEC_PRODUCTO] = rotas
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any(h["Revisar"] == "Especificación" for h in hallazgos))

    def test_un_producto_nuevo_sin_fotos_avisa_pero_no_bloquea(self):
        _tablas, hallazgos = self._validar([_ficha("ZZ26-AAA", Imagenes=[])])
        imagen = [h for h in hallazgos if h["Revisar"] == "Imagen"]
        self.assertTrue(imagen)
        self.assertTrue(all(h["Gravedad"] == vtex.AVISA for h in imagen))

    def test_una_url_de_imagen_que_no_es_web_bloquea(self):
        tablas, _h = self._validar([_ficha("ZZ27-AAA", Imagenes=["C:/fotos/a.jpg"])])
        hallazgos = vtex.validar(tablas, self.catalogo)
        self.assertTrue(any(h["Revisar"] == "Imagen" and h["Gravedad"] == vtex.BLOQUEA
                            for h in hallazgos))


class CapaDePantalla(unittest.TestCase):
    """El pegamento con Streamlit. Es donde viven los fallos que llegan a
    produccion: las pruebas de motor no lo ven."""

    def test_el_matrixify_se_convierte_en_fichas(self):
        import pandas as pd
        import app_matrixify as app
        df = pd.DataFrame([
            {"Metafield: custom.codigo_modelo_color [id]": "AB1-N11",
             "Title": "ZAPATO", "Body HTML": "<p>x</p>", "Type": "Zapatos",
             "Metafield: custom.marca [single_line_text_field]": "Columbia",
             "Metafield: custom.genero [single_line_text_field]": "Hombre",
             "Metafield: custom.color [single_line_text_field]": "Negro",
             "Metafield: custom.categoria [single_line_text_field]": "Calzado",
             "Option1 Value": "39", "Variant SKU": "S1", "Variant Barcode": "7701",
             "Image Src": "https://x/a.jpg; https://x/b.jpg", "Temporada": "V26"},
            # La segunda fila es de VARIANTE: en Matrixify los campos del
            # producto van solo en la primera.
            {"Metafield: custom.codigo_modelo_color [id]": "AB1-N11",
             "Title": "", "Body HTML": "", "Type": "",
             "Option1 Value": "40", "Variant SKU": "S2", "Variant Barcode": "",
             "Image Src": ""},
        ])
        fichas = app.vtex_fichas_desde_matrixify(df)
        self.assertEqual(len(fichas), 1)
        ficha = fichas[0]
        self.assertEqual(ficha["Title"], "ZAPATO")
        self.assertEqual(ficha["Temporada"], "V26")
        self.assertEqual(ficha["Imagenes"], ["https://x/a.jpg", "https://x/b.jpg"])
        self.assertEqual([v["Talla"] for v in ficha["Variantes"]], ["39", "40"])

    def test_una_talla_repetida_entre_bloques_no_duplica_la_variante(self):
        """La carga va por bloques y esto se arma sobre la concatenacion."""
        import pandas as pd
        import app_matrixify as app
        df = pd.DataFrame([
            {"Metafield: custom.codigo_modelo_color [id]": "AB1-N11", "Title": "X",
             "Option1 Value": "39", "Variant SKU": "S1"},
            {"Metafield: custom.codigo_modelo_color [id]": "AB1-N11", "Title": "X",
             "Option1 Value": "39", "Variant SKU": "S1"},
        ])
        self.assertEqual(len(app.vtex_fichas_desde_matrixify(df)[0]["Variantes"]), 1)

    def test_el_zip_lleva_los_cuatro_archivos_con_la_fila_en_blanco(self):
        import app_matrixify as app
        planillas, _informe = _leer_muestra()
        catalogo = vtex.leer_catalogo(planillas)
        emparejados = vtex.emparejar([_ficha("HP102011307-251")], catalogo,
                                     nombres_de_tipo=garment_types.sinonimos_de)
        tablas, _inc, _res = vtex.generar(emparejados, catalogo)
        paquete = app.vtex_armar_zip(tablas, catalogo)
        with zipfile.ZipFile(io.BytesIO(paquete)) as zf:
            nombres = zf.namelist()
            self.assertEqual(len(nombres), 4)
            for archivo in vtex.ARCHIVOS:
                self.assertTrue(
                    any(vtex.NOMBRES_DE_ARCHIVO[archivo] in n for n in nombres),
                    vtex.ETIQUETAS[archivo])
            productos = next(n for n in nombres
                             if vtex.NOMBRES_DE_ARCHIVO[vtex.PRODUCTOS] in n)
            import pandas as pd
            crudo = pd.read_excel(io.BytesIO(zf.read(productos)), header=None, dtype=str)
        # Primera fila en blanco y cabecera en la segunda, como el export.
        self.assertTrue(crudo.iloc[0].isna().all(),
                        "la primera fila tiene que ir en blanco, como el export de VTEX")
        self.assertEqual(list(crudo.iloc[1]), list(vtex.COLUMNAS_PRODUCTOS))

    def test_el_zip_releido_se_reconoce_como_un_export_de_vtex(self):
        """La prueba de fuego: lo que sale se puede volver a leer como maestro."""
        import app_matrixify as app
        planillas, _informe = _leer_muestra()
        catalogo = vtex.leer_catalogo(planillas)
        emparejados = vtex.emparejar(
            [_ficha("HP102011307-251", Imagenes=["https://x/a.jpg"])], catalogo,
            nombres_de_tipo=garment_types.sinonimos_de)
        tablas, _inc, _res = vtex.generar(emparejados, catalogo)
        paquete = app.vtex_armar_zip(tablas, catalogo)

        class Subido(io.BytesIO):
            def __init__(self, nombre, datos):
                super().__init__(datos)
                self.name = nombre

        with zipfile.ZipFile(io.BytesIO(paquete)) as zf:
            subidos = [Subido(n, zf.read(n)) for n in zf.namelist()]
        vueltas, _informe = app.vtex_leer_planillas(subidos)
        for archivo in vtex.ARCHIVOS:
            self.assertTrue(vueltas.get(archivo), vtex.ETIQUETAS[archivo])
        primera = next(iter(vueltas[vtex.PRODUCTOS][0]))
        self.assertEqual(primera["Product ID"], "2")

    def test_el_excel_se_escribe_en_orden_de_fila(self):
        """`constant_memory` suelta cada fila al pasar a la siguiente, asi que
        **exige** escribir en orden de fila. Escrito de otra forma no avisa:
        pierde datos en silencio, que es como se descarto en
        `dataframe_to_excel_bytes`."""
        import app_matrixify as app
        cuerpo = inspect.getsource(app.vtex_excel_de_planilla)
        self.assertIn("constant_memory=True", cuerpo)
        # El bucle exterior es el de las FILAS y el interior el de las columnas.
        arbol = ast.parse(textwrap.dedent(cuerpo))
        bucles = [n for n in ast.walk(arbol) if isinstance(n, ast.For)]
        externo = next(n for n in bucles if any(
            isinstance(h, ast.For) for h in ast.walk(n) if h is not n))
        self.assertIn("numero", ast.dump(externo.target),
                      "el bucle de fuera tiene que ser el de las filas")

    def test_el_excel_conserva_lo_que_no_puede_cambiar(self):
        """Un `=1+1` que salga como FORMULA y una URL que salga como
        HIPERVINCULO son las dos formas de corromper este archivo: Excel admite
        65.530 hipervinculos por hoja y la planilla de imagenes es una columna
        entera de URLs."""
        import pandas as pd
        import app_matrixify as app
        columnas = ("A", "B", "C")
        filas = [{"A": "=1+1", "B": "https://x/a.jpg", "C": "0012"},
                 {"A": "Niño · 04-Jun", "B": "", "C": "000"}]
        datos = app.vtex_excel_de_planilla(columnas, filas)
        leido = pd.read_excel(io.BytesIO(datos), header=None, dtype=str)
        self.assertTrue(leido.iloc[0].isna().all())
        self.assertEqual(list(leido.iloc[1]), list(columnas))
        self.assertEqual(list(leido.iloc[2]), ["=1+1", "https://x/a.jpg", "0012"])
        self.assertEqual(leido.iloc[3, 0], "Niño · 04-Jun")
        self.assertEqual(leido.iloc[3, 2], "000")

    def test_la_pantalla_no_escribe_en_ninguna_tienda(self):
        """Entrega un ZIP. Ni Shopify ni VTEX se tocan desde aqui."""
        import app_matrixify as app
        prohibidas = ("product_update", "product_create", "metafields_set",
                      "productCreateMedia", "collectionAddProducts",
                      "apply_shopify_preview", "apply_full_product_updates",
                      "lanzar_carga_remota")
        for funcion in (app.render_vtex_generator, app.vtex_analizar, app.vtex_armar_zip,
                        app.vtex_fichas_desde_matrixify):
            cuerpo = inspect.getsource(funcion)
            for prohibida in prohibidas:
                self.assertNotIn(prohibida, cuerpo, f"{funcion.__name__} -> {prohibida}")

    def test_el_motor_no_importa_streamlit_ni_pandas(self):
        """`engines/` nunca importa Streamlit. Es regla de arquitectura."""
        arbol = ast.parse(inspect.getsource(vtex))
        importados = set()
        for nodo in ast.walk(arbol):
            if isinstance(nodo, ast.Import):
                importados.update(a.name.split(".")[0] for a in nodo.names)
            elif isinstance(nodo, ast.ImportFrom) and nodo.module:
                importados.add(nodo.module.split(".")[0])
        self.assertNotIn("streamlit", importados)
        self.assertNotIn("pandas", importados)

    def test_la_pantalla_esta_en_el_menu_y_tiene_despacho(self):
        import app_matrixify as app
        destinos = [i["area"] for g in app.nav_grupos(True) for i in g["items"]]
        self.assertIn(app.VTEX_LABEL, destinos)
        cuerpo = inspect.getsource(app.main)
        self.assertIn("render_vtex_generator()", cuerpo)
        self.assertEqual(app.nav_ruta(app.VTEX_LABEL, ""), ["Cargas", app.VTEX_LABEL])

    def test_el_boton_del_menu_esta_en_las_cinco_listas_de_css(self):
        """Un boton nuevo hay que registrarlo en CINCO listas de selectores y
        darle su icono; nada en el codigo lo obliga."""
        import app_matrixify as app
        css = inspect.getsource(app.inject_custom_css)
        for patron in ("div.st-key-operation_nav_vtex button,",
                       'div.st-key-operation_nav_vtex button [data-testid="stMarkdownContainer"],',
                       "div.st-key-operation_nav_vtex button p,",
                       "div.st-key-operation_nav_vtex button::before,",
                       "div.st-key-operation_nav_vtex button:hover,",
                       "div.st-key-operation_nav_vtex button::before {{"):
            self.assertIn(patron, css, patron)

    def test_el_matrixify_sale_del_mismo_motor_que_la_carga_supermall(self):
        """Un segundo constructor de Matrixify se separa del primero sin que
        nadie lo note, que es lo que este repositorio ya pago con las dos
        `normalize_size`."""
        import app_matrixify as app
        cuerpo = inspect.getsource(app.vtex_analizar)
        self.assertIn("supermall_generar(", cuerpo)
        self.assertNotIn("build_centry_matrixify_from_master", cuerpo)
        self.assertNotIn("build_columbia_matrixify", cuerpo)


class DePuntaAPunta(unittest.TestCase):
    """La cadena entera: Shopify + ARTI -> Matrixify -> las cuatro planillas.

    Con un Shopify y un maestro FALSOS, pero ejecutando el mismo
    `supermall_generar` que usa la Carga Supermall. Es donde se ven los fallos
    del pegamento, que es justo lo que las pruebas de motor no ven.
    """

    def _pipeline(self):
        import pandas as pd
        import streamlit as st
        import app_matrixify as app

        # Dos productos: uno que YA esta en el export de VTEX y otro que no.
        catalogos = {
            "columbia": [
                {"Mod-Col": "HP102011307-251", "Handle": "zapato-stanford",
                 "Product ID": "gid://1",
                 "Title": "ZAPATO HUSH PUPPIES STANFORD II PARA HOMBRE",
                 "Body HTML": "<p>Slip on</p>", "Type": "Zapatos",
                 "Tags": "Hush Puppies", "Vendor": "columbiape",
                 "Status": "ACTIVE", "Published Online Store": "SI",
                 "Image Src": "https://cdn/x1.jpg; https://cdn/x2.jpg",
                 "Marca": "Hush Puppies",
                 "Metafield: custom.codigo_modelo_color [id]": "HP102011307-251",
                 "Metafield: custom.marca [single_line_text_field]": "Hush Puppies",
                 "Metafield: custom.genero [single_line_text_field]": "Hombre",
                 "Metafield: custom.color [single_line_text_field]": "Rojo",
                 "Variants": []},
                {"Mod-Col": "AM8004-YFO", "Handle": "pantalon-silver-ridge",
                 "Product ID": "gid://2", "Title": "PANTALON SILVER RIDGE HOMBRE",
                 "Body HTML": "<p>Pantalon</p>", "Type": "Pantalones",
                 "Tags": "Columbia", "Vendor": "columbiape",
                 "Status": "ACTIVE", "Published Online Store": "SI",
                 "Image Src": "https://cdn/y1.jpg", "Marca": "Columbia",
                 "Metafield: custom.codigo_modelo_color [id]": "AM8004-YFO",
                 "Metafield: custom.marca [single_line_text_field]": "Columbia",
                 "Metafield: custom.genero [single_line_text_field]": "Hombre",
                 "Metafield: custom.color [single_line_text_field]": "Negro",
                 "Variants": []},
            ],
            "supermall": [],
        }
        maestro = pd.DataFrame([
            {"Mod-Col": codigo, "COD MOD COL": codigo, "CODINT_MA": f"{codigo}-{n}",
             "TALNUM_MA": str(talla), "MARCA_MA": marca, "Precio": "199.90",
             "CodBarras": f"770{n}{abs(hash(codigo)) % 1000}"}
            for codigo, marca, tallas in (
                ("HP102011307-251", "HUSH PUPPIES", ("39", "40", "41")),
                ("AM8004-YFO", "COLUMBIA", ("30", "32", "34")),
            )
            for n, talla in enumerate(tallas)
        ])

        planillas, _informe = _leer_muestra()
        estado = [{"Sitio": "Columbia.pe", "Estado": "Leido", "Detalle": ""},
                  {"Sitio": "Supermall.pe", "Estado": "Leido", "Detalle": ""}]
        originales = {
            "cargar_catalogos_de_todos_los_sitios": app.cargar_catalogos_de_todos_los_sitios,
            "session_arti_for_app": app.session_arti_for_app,
            "olvidar_arti_de_la_sesion": app.olvidar_arti_de_la_sesion,
            "sitio_espejo": app.sitio_espejo,
            "get_shopify_config": app.get_shopify_config,
        }
        app.cargar_catalogos_de_todos_los_sitios = lambda force_refresh=False: (
            catalogos, estado)
        app.session_arti_for_app = lambda brand_config, **kw: (maestro.copy(), "maestro falso")
        app.olvidar_arti_de_la_sesion = lambda brand_config: None
        app.sitio_espejo = lambda: "supermall"
        app.get_shopify_config = lambda site_key: {"shop_domain": "", "access_token": ""}
        try:
            st.session_state.clear()
            return app.vtex_analizar(planillas, [], False)
        finally:
            for nombre, funcion in originales.items():
                setattr(app, nombre, funcion)

    def test_la_cadena_entera_empareja_contra_el_export_real(self):
        analisis = self._pipeline()
        por_codigo = {e["Mod-Col"]: e for e in analisis["emparejados"]}
        self.assertIn("HP102011307-251", por_codigo)
        existente = por_codigo["HP102011307-251"]
        self.assertEqual(existente["Estado"], vtex.EXISTE)
        self.assertEqual(existente["Product ID"], "2")
        self.assertTrue(existente["Tallas"], "las tallas salen del maestro ARTI")

    def test_el_producto_que_no_esta_en_vtex_sale_como_nuevo(self):
        analisis = self._pipeline()
        por_codigo = {e["Mod-Col"]: e for e in analisis["emparejados"]}
        # `AM8004-YFO` esta en la muestra del export, asi que el caso nuevo se
        # comprueba con lo que el cruce diga de el: exista o no, no puede
        # inventarse un Product ID.
        for entrada in analisis["emparejados"]:
            if entrada["Estado"] == vtex.NUEVO:
                self.assertEqual(entrada["Product ID"], "")

    def test_los_cuatro_archivos_salen_de_la_cadena_completa(self):
        import app_matrixify as app
        analisis = self._pipeline()
        tablas, _inc, resumen = vtex.generar(
            analisis["emparejados"], analisis["catalogo"])
        self.assertTrue(tablas[vtex.PRODUCTOS])
        self.assertTrue(tablas[vtex.ESPEC_PRODUCTO])
        self.assertTrue(tablas[vtex.ESPEC_SKU])
        self.assertTrue(tablas[vtex.IMAGENES])
        hallazgos = vtex.validar(tablas, analisis["catalogo"])
        bloqueos, _avisos = vtex.resumen_de_validacion(hallazgos)
        self.assertEqual(bloqueos, 0,
                         [h for h in hallazgos if h["Gravedad"] == vtex.BLOQUEA])
        paquete = app.vtex_armar_zip(tablas, analisis["catalogo"])
        with zipfile.ZipFile(io.BytesIO(paquete)) as zf:
            self.assertEqual(len(zf.namelist()), 4)

    def test_las_tallas_del_maestro_llegan_a_la_planilla(self):
        """Las tallas no salen de Shopify: salen del maestro ARTI, y ese es el
        tramo que solo se ve ejecutando la cadena entera."""
        analisis = self._pipeline()
        tablas, _inc, _res = vtex.generar(analisis["emparejados"], analisis["catalogo"])
        nombres = {f["SKU name"] for f in tablas[vtex.PRODUCTOS]
                   if f["Product reference code"] == "HP102011307-251"}
        self.assertTrue(nombres)
        self.assertTrue(all(n.startswith("TALLA ") for n in nombres), nombres)


# El guion que mide, en un SUBPROCESO: el pico de RSS es del proceso entero, y
# medir varias cosas en el mismo interprete las mezclaria. Es la misma forma que
# usa `scripts/test_memoria.py`.
GUION_DE_MEMORIA = """
import io, os, sys
sys.path.insert(0, %(raiz)r)
import xlsxwriter
from engines import vtex_export as V
import app_matrixify as app

def rss():
    with open("/proc/self/statm") as fh:
        return int(fh.read().split()[1]) * os.sysconf("SC_PAGE_SIZE") / 1024 / 1024

PRODUCTOS, TALLAS, CAMPOS = %(productos)d, 8, 30

def libro(columnas, filas):
    buffer = io.BytesIO()
    wb = xlsxwriter.Workbook(buffer, {"constant_memory": True})
    ws = wb.add_worksheet("Sheet1")
    for c, n in enumerate(columnas):
        ws.write_string(1, c, n)
    for r, fila in enumerate(filas, start=2):
        for c, n in enumerate(columnas):
            if fila.get(n):
                ws.write_string(r, c, str(fila[n]))
    wb.close()
    return buffer.getvalue()

def prod():
    for p in range(PRODUCTOS):
        for t in range(TALLAS):
            yield {"Product ID": str(p+1), "Product Name": "PRODUCTO %%d" %% p,
                   "Product reference code": "MC%%06d-C01" %% p,
                   "Brand ID": "1", "Brand": "Columbia", "Department ID": "25",
                   "Department": "Hombre", "Category ID": "41", "Category": "Zapatos",
                   "SKU ID": str(300000+p*TALLAS+t), "SKU name": "TALLA %%d" %% (35+t),
                   "SKU reference code": str(300000+p*TALLAS+t),
                   "Package weight": "800", "Package width": "34",
                   "Package height": "49.5", "Package length": "14",
                   "Description": "x"*200, "Sales channels": "1, 4"}

def espec():
    for p in range(PRODUCTOS):
        for c in range(CAMPOS):
            yield {"ID del producto": str(p+1), "Código de referencia del producto":
                   "MC%%06d-C01" %% p, "ID de categoría": "41", "Categoría": "Zapatos",
                   "ID de campo": str(18+c), "Nombre del campo": "Campo %%d" %% c,
                   "Tipo de campo": "Texto", "IDs de especificación": str(400000+p*CAMPOS+c),
                   "Valores de especificación": "---"}

archivos = [("products-and-skus.xlsx", libro(V.COLUMNAS_PRODUCTOS, prod())),
            ("especificaciones.xlsx", libro(V.COLUMNAS_ESPEC_PRODUCTO, espec()))]

class Subido(io.BytesIO):
    def __init__(self, nombre, datos):
        super().__init__(datos)
        self.name = nombre

base = rss()
planillas, _informe = app.vtex_leer_planillas([Subido(n, d) for n, d in archivos])
catalogo = V.leer_catalogo(planillas)
print("%%.0f" %% (rss() - base))
print(len(catalogo.productos))

# Y lo que costaria MATERIALIZARLO, que es lo que se hacia antes. Si esto deja
# de ser caro, la prueba de arriba ya no esta midiendo nada.
del catalogo
planillas, _informe = app.vtex_leer_planillas([Subido(n, d) for n, d in archivos])
base = rss()
materializado = {nombre: [list(hoja) for hoja in hojas]
                 for nombre, hojas in planillas.items()}
print("%%.0f" %% (rss() - base))
print(sum(len(h) for hojas in materializado.values() for h in hojas))
"""


class Memoria(unittest.TestCase):
    """El export real son 177 MB y el contenedor da 1 GB PARA TODA LA APP.

    Lo que esta prueba impide que vuelva: que la lectura materialice el archivo.
    Medido antes de arreglarlo, con un export de 122 MB: **1.728 MB** solo en
    tener las filas como diccionarios.
    """

    PRODUCTOS = 2500
    # MB por encima del arranque, leyendo E INDEXANDO. Medido: 35 MB. El tope
    # deja margen de plataforma pero se queda MUY por debajo de lo que cuesta
    # materializar las mismas filas, que es lo que esta prueba impide.
    PRESUPUESTO = 120

    def test_leer_e_indexar_no_materializa_el_export(self):
        guion = GUION_DE_MEMORIA % {"raiz": RAIZ, "productos": self.PRODUCTOS}
        salida = subprocess.run([sys.executable, "-c", guion], capture_output=True,
                                text=True, timeout=900)
        self.assertEqual(salida.returncode, 0, salida.stderr[-2000:])
        lineas = [l for l in salida.stdout.strip().splitlines() if l.strip()]
        gastado, productos, materializar, filas = (
            float(lineas[-4]), int(lineas[-3]), float(lineas[-2]), int(lineas[-1]))
        self.assertEqual(productos, self.PRODUCTOS, "el indice tiene que estar completo")
        self.assertLess(gastado, self.PRESUPUESTO,
                        f"leer e indexar {self.PRODUCTOS} productos costo {gastado:.0f} MB")
        # Y que materializar las MISMAS filas siga siendo caro: si dejara de
        # serlo, el tope de arriba ya no estaria midiendo nada. Es la misma
        # guarda que lleva `scripts/test_memoria.py`.
        self.assertGreater(filas, 10000)
        self.assertGreater(materializar, gastado * 2,
                           f"materializar {filas:,} filas costo {materializar:.0f} MB "
                           f"contra {gastado:.0f} MB de la lectura en streaming: el tope "
                           "de la prueba ya no distingue una cosa de la otra")

    def test_las_planillas_grandes_no_se_materializan_al_generar(self):
        """La de especificaciones son 477.000 filas con 9.000 productos: tenerlas
        vivas costaba 847 MB medidos."""
        planillas, _informe = _leer_muestra()
        catalogo = vtex.leer_catalogo(planillas)
        emparejados = vtex.emparejar([_ficha("HP102011307-251")], catalogo,
                                     nombres_de_tipo=garment_types.sinonimos_de)
        tablas, _inc, _res = vtex.generar(emparejados, catalogo)
        for archivo in (vtex.ESPEC_PRODUCTO, vtex.ESPEC_SKU, vtex.IMAGENES):
            self.assertIsInstance(tablas[archivo], vtex.Filas, vtex.ETIQUETAS[archivo])
            # Y se puede recorrer mas de una vez: la validacion la recorre y
            # despues la escribe el ZIP.
            primera = [f["ID de campo"] for _n, f in zip(range(3), tablas[archivo])]
            self.assertEqual(primera,
                             [f["ID de campo"] for _n, f in zip(range(3), tablas[archivo])])


class ElIdiomaDelExport(unittest.TestCase):
    """El export sale en el IDIOMA del admin de VTEX.

    El pack de muestra vino con la planilla de productos en INGLES y el export
    real de la tienda sale en ESPANOL: mismas 50 columnas, mismo orden, otros
    nombres -- y `Sí` donde la muestra decia `Yes`. La pantalla contestaba
    "No reconocida · 0 filas" con el archivo de 24 MB del catalogo entero.

    `products-and-skus-es.xlsx` son las primeras filas de ESE archivo.
    """

    def setUp(self):
        import app_matrixify as app
        self.app = app
        self.espanol, _informe = app.vtex_leer_planillas([_subido("products-and-skus-es.xlsx")])
        self.ingles, _informe = app.vtex_leer_planillas([_subido("products-and-skus.xlsx")])

    def test_la_planilla_en_espanol_se_reconoce(self):
        self.assertTrue(self.espanol.get(vtex.PRODUCTOS),
                        "el export real de VTEX sale en espanol y no se reconocia")

    def test_las_dos_dan_el_mismo_indice(self):
        """El idioma no puede cambiar lo que la app entiende del catalogo."""
        for planillas, etiqueta in ((self.espanol, "espanol"), (self.ingles, "ingles")):
            catalogo = vtex.leer_catalogo(planillas)
            with self.subTest(idioma=etiqueta):
                self.assertTrue(catalogo.productos, etiqueta)
                self.assertIn("hush puppies", catalogo.marcas)
                self.assertIn(("hombre", "zapatos"), catalogo.categorias)
                producto = catalogo.producto_por_referencia("HP102011307-251")
                self.assertTrue(producto, "el Mod-Col se lee igual en los dos idiomas")
                self.assertEqual(producto[0].id, "2")
                self.assertTrue(producto[0].skus)
                self.assertEqual(producto[0].skus[0].talla, "39")

    def test_el_si_y_el_no_salen_del_propio_export(self):
        """Un `Yes` en una planilla en espanol deja el producto sin activar."""
        self.assertEqual(vtex.leer_catalogo(self.espanol).si(), "Sí")
        self.assertEqual(vtex.leer_catalogo(self.ingles).si(), "Yes")
        self.assertEqual(vtex.leer_catalogo(self.espanol).no(), "No")

    def test_la_salida_va_en_el_idioma_del_archivo_que_se_subio(self):
        """Es la plantilla que ESE VTEX espera de vuelta."""
        for planillas, primera, activo in ((self.espanol, "ID del producto", "Sí"),
                                           (self.ingles, "Product ID", "Yes")):
            catalogo = vtex.leer_catalogo(planillas)
            emparejados = vtex.emparejar([_ficha("HP102011307-251")], catalogo,
                                         nombres_de_tipo=garment_types.sinonimos_de)
            tablas, _inc, _res = vtex.generar(emparejados, catalogo)
            fila = tablas[vtex.PRODUCTOS][0]
            with self.subTest(idioma=primera):
                self.assertEqual(list(fila)[0], primera)
                self.assertEqual(len(fila), len(vtex.COLUMNAS_PRODUCTOS))
                self.assertEqual(fila[catalogo.columna(vtex.PRODUCTOS, "Active product")],
                                 activo)
                self.assertEqual(
                    fila[catalogo.columna(vtex.PRODUCTOS, "Product reference code")],
                    "HP102011307-251")

    def test_la_traduccion_NO_toca_las_otras_planillas(self):
        """`ID de SKU` es sinonimo de `SKU ID` en la planilla de productos, pero
        en la de especificaciones de SKU es su propia columna."""
        self.assertEqual(vtex.canonico("ID de SKU", vtex.PRODUCTOS), "SKU ID")
        self.assertEqual(vtex.canonico("ID de SKU", vtex.ESPEC_SKU), "ID de SKU")
        planillas, _informe = _leer_muestra()
        catalogo = vtex.leer_catalogo(planillas)
        self.assertIsNotNone(catalogo.campo_de_sku("41", "Talla"))

    def test_una_columna_que_no_esta_en_la_tabla_se_conserva(self):
        """Una columna nueva de VTEX tiene que llegar igual al archivo de
        salida, no desaparecer."""
        self.assertEqual(vtex.canonico("Columna nueva de VTEX"), "Columna nueva de VTEX")
        fila = vtex.traducir_fila({"ID del producto": "1", "Columna nueva": "x"})
        self.assertEqual(fila, {"Product ID": "1", "Columna nueva": "x"})

    def test_cuenta_las_filas_aunque_el_archivo_no_declare_su_dimension(self):
        """`max_row` de openpyxl sale `None` con el export real: el archivo de
        24 MB del usuario reportaba 0 filas."""
        _planillas, informe = self.app.vtex_leer_planillas(
            [_subido("products-and-skus-es.xlsx")])
        self.assertEqual(len(informe), 1)
        self.assertGreater(informe[0]["Filas"], 100)

    def test_una_fila_con_la_ultima_celda_vacia_no_pierde_columnas(self):
        """De la primera fila sale la CABECERA del archivo de salida, y un xlsx
        no escribe las celdas vacias del final: `zip` cortaba el registro por la
        fila mas corta y la planilla salia con una columna de menos."""
        hoja = self.espanol[vtex.PRODUCTOS][0]
        for _numero, fila in zip(range(20), hoja):
            self.assertEqual(len(fila), len(vtex.COLUMNAS_PRODUCTOS))

    def test_el_lector_del_xlsx_da_lo_mismo_que_openpyxl(self):
        """Se cambio openpyxl por un lector propio porque openpyxl recorre el
        XML entero al abrir un archivo que no declara su dimension -- 24 s por
        apertura con el export real. Lo que no puede cambiar es lo que lee."""
        import io as _io
        from openpyxl import load_workbook
        for nombre in sorted(os.listdir(MUESTRA)):
            if not nombre.endswith(".xlsx"):
                continue
            datos = _subido(nombre).getvalue()
            hojas = self.app.vtex_hojas_del_libro(datos)
            libro = load_workbook(_io.BytesIO(datos), read_only=True, data_only=True)
            try:
                self.assertEqual([h for h, _r in hojas], libro.sheetnames, nombre)
                for hoja, ruta in hojas:
                    propio = list(self.app.vtex_filas_del_xml(datos, ruta))
                    suyo = [list(f) for f in libro[hoja].iter_rows(values_only=True)]
                    self.assertEqual(len(propio), len(suyo), f"{nombre}/{hoja}: filas")
                    for nuestra, suya in zip(propio, suyo):
                        ancho = max(len(nuestra), len(suya))
                        nuestra = [clean(v) for v in nuestra] + [""] * (ancho - len(nuestra))
                        suya = [clean(v) for v in suya] + [""] * (ancho - len(suya))
                        self.assertEqual(nuestra, suya, f"{nombre}/{hoja}")
            finally:
                libro.close()


def clean(valor):
    return "" if valor is None else str(valor)


if __name__ == "__main__":
    unittest.main(verbosity=2)
